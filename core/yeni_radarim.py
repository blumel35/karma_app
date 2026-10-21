# core/yeni_radarim.py
# -*- coding: utf-8 -*-
"""
YENİ Radarım — FAZ 1: "Portföylerim" akışı (mockup v13).

Yalnızca yeni ana sayfanın (core/yeni_ana_sayfa.py) "Radarım" sayfasında,
özellik anahtarı arkasında çalışır. Mevcut Radarım (pages/Danisman_Bolgem.py,
core/bolgem.py) ve danışmanların ekranı DEĞİŞMEDİ. SALT OKUNUR, bildirim yok.

Akış A:  Portföylerim  →  [ Alıcı talepleri | Benzer Startkey portföyleri | Benzer FSBO ilanları ]  →  Eşleşmeler
Akış B:  Taleplerim  →  [ Zeta portföyleri | Startkey portföyleri | FSBO ilanları | Portföy paylaşımları ]  →  Eşleşmeler
  (B'deki portföy kutuları, taleplerimden en az biriyle OLASI eşiğini (%40) geçen portföyleri sayar;
   Eşleşmeler = güçlü (%70+) talep→portföy çiftleri. Bir talep birden çok ilçe istediyse her birinde sayılır.)

Yeni puan/eşleşme kuralı YOK:
  - Veri: core.eslesme_havuzu.veri_yukle() (Bölge Havuzu "Eşleşme" sekmesiyle aynı).
  - Eşleşme: core.eslesme_havuzu.portfoyden_taleplere() (ilçe 40 / bütçe 35 /
    oda 15 / kriter 10; işlem-mülk tipi uyumsuzsa eleme; eşik güçlü 70, olası 40).
  - Talep/paylaşım ömrü: es.omre_gore (kiralık 14 / satılık 45 gün), Zeta/Startkey aktif stok.
  - Bayat (Revy'de görünmeyen) Zeta ilanları veri_yukle'de zaten dışarıda.
Bölge kapsamı, mevcut Radarım ile aynı karar (12.10.2026): Uzmanlık Bölgelerim (en fazla 5)
+ kendi kayıtlarının bulunduğu ilçeler.

"Benzer Startkey portföyü" (bu sayfada tanımlı, motorun yardımcılarıyla): aynı ilçe,
aynı işlem ve mülk tipi, fiyat ±%30 içinde, oda sayısı en fazla 1 fark. Fiyatı okunamayan
ilan "benzer" sayılmaz (iddiayı sağlam tutmak için).
"""
import html as _html
import re
from datetime import datetime, timedelta, timezone

import streamlit as st

from core import bolge_havuzu as bh
from core import bolgem as bg
from core import eslesme_havuzu as es
from core.eslesme_kart_bilesen import eslesme_kartlari
from core.havuz_tablo_bilesen import havuz_tablosu
from core.bolge_secici import bolgelerini_cek
from core.danisman_ortak import su_anki_danisman, su_anki_danisman_varyantlari, IZMIR_ILCELERI
from core.match_engine import (
    _mulk_normalize, _islem_normalize, _belirsiz_mi, _sayi_ayikla, _oda_secenekleri_cikar,
)

# "Portföy paylaşımı" = ofis içinde ELLE girilen (kapalı) portföyler — Zeta Paylaşımları sayfasıyla aynı kural
# (kaynak: zeta / ofis). Mail sistemi, Startkey vb. kaynaklı kayıtlar paylaşım SAYILMAZ.
PAYLASIM_KAYNAKLARI = {"zeta", "ofis"}
# Yönetici "Tüm ekip" görünümünde eşleşme hesabı en yeni bu kadar kaynak kayıt için yapılır (süre sınırı).
EKIP = "👥 Tüm ekip"
PIYASA = "🌐 Tüm Startkey"
EKIP_PORT_SINIR = 300
EKIP_TALEP_SINIR = 200
YONETICILER = {"meltem bulu"}
MAX_BOLGE = 5
BOLGE_TABLOLARI = ("uzmanlik_bolgeleri", "fsbo_bolgeleri", "startkey_ilan_bolgeleri")
TUM = "Tüm uzmanlık bölgelerim"
BENZER_FIYAT_ORAN = 0.30
GRUP_SAYFA = 15


# ── VERİ ────────────────────────────────────────────────────────────────
@st.cache_data(ttl=300, show_spinner="Radarım verisi yükleniyor...")
def _veri_temel():
    """Alıcı talepleri + Zeta portföyleri + paylaşımlar (core.eslesme_havuzu.veri_yukle ile aynı kurallar).
    Startkey/FSBO burada YOK: Bölge Havuzu okuyucusu tablo başına en son 3000 kayıtla sınırlı olduğu için
    bu iki stok ayrıca, ilçe ilçe okunur (_pazar_stok)."""
    simdi = datetime.now(timezone.utc)
    kesim = simdi - timedelta(days=es.ADAY_PENCERESI_GUN)
    veri = {"talep": [], "zeta": [], "paylasim": []}
    hatalar, bilgi = {}, {"bayat_zeta": 0}
    try:
        veri["talep"] = [es._talep_kaydi(k) for k in bh._alici_talepleri(kesim, simdi)]
    except Exception as e:      # noqa: BLE001
        hatalar["Alıcı talebi"] = str(e)
    try:
        paylasim, zeta = bh._portfoyler(kesim, simdi)
        veri["paylasim"] = [es._portfoy_kaydi(k, "Paylaşım") for k in paylasim]
        tum_zeta = [es._portfoy_kaydi(k, "Zeta portföyü") for k in zeta]
        veri["zeta"] = [k for k in tum_zeta if not k["bayat"]]      # Revy'de görünmeyen (bayat) ilanlar eşleşmeye girmez
        bilgi["bayat_zeta"] = len(tum_zeta) - len(veri["zeta"])
    except Exception as e:      # noqa: BLE001
        hatalar["Portföyler"] = str(e)
    return veri, hatalar, bilgi


@st.cache_data(ttl=900, show_spinner="Startkey ve FSBO ilanları ilçe ilçe okunuyor...")
def _pazar_stok():
    """Startkey ve FSBO aktif stoku — 3000 kayıt sınırına takılmamak için her İzmir ilçesi ayrı sorgulanır
    (paralel). Dönüş: ({'startkey': [...], 'fsbo': [...]}, hatalar, sınıra_takılan_ilçeler)."""
    from concurrent.futures import ThreadPoolExecutor
    simdi = datetime.now(timezone.utc)
    kesim = simdi - timedelta(days=3650)
    gorevler = [(tur, ilce) for tur in ("startkey", "fsbo") for ilce in IZMIR_ILCELERI]

    def _oku(g):
        tur, ilce = g
        try:
            return g, (bh._startkey if tur == "startkey" else bh._fsbo)(kesim, simdi, [ilce]), ""
        except Exception as e:      # noqa: BLE001
            return g, [], str(e)

    sonuc, hatalar, sinir = {"startkey": [], "fsbo": []}, {}, []
    etiket = {"startkey": "Startkey ilanı", "fsbo": "FSBO ilanı"}
    with ThreadPoolExecutor(max_workers=6) as ex:
        for (tur, ilce), kayitlar, hata in ex.map(_oku, gorevler):
            if hata:
                hatalar[f"{etiket[tur]} ({ilce})"] = hata
                continue
            if len(kayitlar) >= bh.MAX_SATIR:
                sinir.append(f"{etiket[tur]}: {ilce}")
            sonuc[tur].extend(es._portfoy_kaydi(k, etiket[tur]) for k in kayitlar)
    return sonuc, hatalar, sinir


def _veri():
    """es.veri_yukle() ile aynı biçim: (veri, hatalar, bilgi) — Startkey/FSBO sınırsız."""
    veri, hatalar, bilgi = _veri_temel()
    pazar, h2, sinir = _pazar_stok()
    veri = dict(veri)
    veri["startkey"], veri["fsbo"] = pazar["startkey"], pazar["fsbo"]
    bilgi = dict(bilgi)
    bilgi["sinir"] = sinir
    return veri, {**hatalar, **h2}, bilgi


def _veri_temizle():
    for f in (_veri_temel, _pazar_stok, _pazar_hafif, _hesapla, _piyasa):
        f.clear()


@st.cache_data(ttl=60, show_spinner=False)
def _bolgeler(kullanici):
    """Danışmanın en fazla 5 bölgesi: Uzmanlık Bölgelerim; hiç yoksa FSBO, sonra Startkey listesi
    (core/bolgem'in mevcut Radarım'daki kuralıyla aynı)."""
    for tablo in BOLGE_TABLOLARI:
        sonuc, gorulen = [], set()
        try:
            for r in bolgelerini_cek(tablo, kullanici):
                i = (r.get("ilce") or "").strip()
                if i and bh.tr_kucuk(i) not in gorulen:
                    gorulen.add(bh.tr_kucuk(i))
                    sonuc.append(i)
        except Exception:      # noqa: BLE001
            continue
        if sonuc:
            return sonuc[:MAX_BOLGE]
    return []


@st.cache_data(ttl=3600, show_spinner=False)
def _zeta_kimlikleri():
    """Zeta personelinin kimlikleri: ({ad_anahtarı: ad}, {e-posta: ad}) — personel listesinden."""
    adlar, epostalar = {}, {}
    try:
        from core.personel_manager import load_personel_listesi
        df = load_personel_listesi()
        for _, r in df.iterrows():
            ad = str(r.get("ad_soyad") or "").strip()
            if not ad:
                continue
            adlar[bh.tr_kucuk(ad)] = ad
            e = str(r.get("email") or "").strip().lower()
            if e and e != "nan":
                epostalar[e] = ad
    except Exception:      # noqa: BLE001
        pass
    return adlar, epostalar


def _sahip_ayir(sahip):
    """'Ad Soyad <eposta>' → (ad, eposta); yalnız e-posta ya da yalnız ad da olabilir."""
    s = str(sahip or "").strip()
    m = re.match(r"^(.*?)\s*<([^>]+)>\s*$", s)
    if m:
        return m.group(1).strip(), m.group(2).strip().lower()
    return ("", s.lower()) if "@" in s else (s, "")


def _zeta_kisisi(sahip):
    """Sahip alanı bir Zeta danışmanına aitse onun adı, değilse ''."""
    adlar, epostalar = _zeta_kimlikleri()
    ad, e = _sahip_ayir(sahip)
    if e and e in epostalar:
        return epostalar[e]
    if ad and bh.tr_kucuk(ad) in adlar:
        return adlar[bh.tr_kucuk(ad)]
    return ""


def _zeta_talebi_mi(k):
    """Radar'da (uygulamadan, kaynak zeta/ofis) girilen ya da Startkey mail sisteminde bir Zeta danışmanı
    tarafından paylaşılan talep. Diğer Startkey ofislerinin mailleri DEĞİL."""
    if str((k.get("ham") or {}).get("kaynak") or "").strip().lower() in PAYLASIM_KAYNAKLARI:
        return True
    return bool(_zeta_kisisi(k.get("sahip")))


def _benim_talebim(k, adlar):
    """bg.benim_mi + mail biçimi: 'Ad Soyad <eposta>' sahibi, e-posta ya da ad varyantlarımla eşleşiyorsa."""
    if bg.benim_mi(k, adlar):
        return True
    ad, e = _sahip_ayir(k.get("sahip"))
    varyant = {bh.tr_kucuk(a) for a in adlar if a}
    vk = {a.lower() for a in adlar if a}
    return bool((ad and bh.tr_kucuk(ad) in varyant) or (e and (e in vk or e.split("@")[0] in vk)))


def _yonetici():
    return str(st.session_state.get("user_name") or "").strip().lower() in YONETICILER


def _yatirim_kaydi(r):
    """Yatırım talebini (musteri_talepleri) motorun alıcı talebi biçimine çevirir; alıcı talepleriyle aynı
    listede/eşleşmede görünür. Bütçe alanının son parçası (en çok) bütçe sayılır."""
    alan = {}
    for a in r.get("alanlar") or []:
        try:
            alan[str(a[0])] = a[1]
        except Exception:      # noqa: BLE001
            continue
    mulk = alan.get("Mülk türü") or ""
    maxb = str(r.get("fiyat") or "").split("–")[-1].strip()
    z = r.get("zaman")
    ham = {
        "talep_eden_danisan": r.get("sahip") or "", "kayit_tarihi": z.isoformat() if z else "",
        "mulk_tipi": mulk, "islem_tipi": "Satılık", "il": "İzmir", "ilce": r.get("ilce") or "",
        "max_butce": maxb, "oda_sayisi_m2": "", "ozel_kriterler": "", "ozet": r.get("alt") or "",
        "bolge_mahalle": r.get("mahalle") or "", "kaynak": "Yatırım talebi",
    }
    kayit = es._talep_kaydi({
        "id": r["id"], "ham": ham, "baslik": "💼 " + (r.get("baslik") or "Yatırım talebi"), "alt": r.get("alt") or "",
        "sahip": r.get("sahip") or "", "kaynak": "Yatırım talebi", "fiyat": r.get("fiyat") or "",
        "ilceler": r.get("ilceler") or [], "zaman": z, "yas_kat": r.get("yas_kat") or "", "islem": "Satılık",
    })
    kayit["yatirim"] = True
    return kayit


@st.cache_data(ttl=300, show_spinner=False)
def _yatirim(adlar, hepsi=False):
    """Yatırım talepleri (kişisel veri): `adlar` verilirse YALNIZ o kişinin kayıtları sunucu tarafında çekilir.
    `hepsi=True` yalnız yönetici "Tüm ekip" görünümünde kullanılır. Dönüş: (kayıtlar, hata_metni)."""
    if not hepsi and not adlar:
        return [], ""
    simdi = datetime.now(timezone.utc)
    try:
        r = bh._musteri_formlari(
            simdi - timedelta(days=3650), simdi, list(IZMIR_ILCELERI),
            danisman_adlari=None if hepsi else list(adlar),
        )
        return [_yatirim_kaydi(k) for k in r.get("yatirim_alici", [])], ""
    except Exception as e:      # noqa: BLE001
        return [], str(e)


def _paylasim_mi(k):
    return str((k.get("ham") or {}).get("kaynak") or "").strip().lower() in PAYLASIM_KAYNAKLARI


def _fiyat(motor):
    return _sayi_ayikla(motor.get("fiyat"))


def _benzer_mi(p, s):
    """p: benim portföyümün motor sözlüğü, s: Startkey ilanının motor sözlüğü."""
    if (p.get("ilce") or "").strip().lower() != (s.get("ilce") or "").strip().lower():
        return False
    a, b = _islem_normalize(p.get("islem_tipi")), _islem_normalize(s.get("islem_tipi"))
    if not _belirsiz_mi(a) and not _belirsiz_mi(b) and a != b:
        return False
    a, b = _mulk_normalize(p.get("mulk_tipi")), _mulk_normalize(s.get("mulk_tipi"))
    if not _belirsiz_mi(a) and not _belirsiz_mi(b) and a != b:
        return False
    fp, fs = _fiyat(p), _fiyat(s)
    if not fp or not fs or abs(fs - fp) > fp * BENZER_FIYAT_ORAN:
        return False
    op, os_ = _oda_secenekleri_cikar(p.get("oda_sayisi_m2")), _oda_secenekleri_cikar(s.get("oda_sayisi_m2"))
    if op and os_ and min(abs(x - y) for x in op for y in os_) > 1:
        return False
    return True


def _bolgede(k, kume):
    """kume None = süzgeç yok (yönetici 'Tüm ilçeler')."""
    return True if kume is None else bg.bolgede_mi(k, kume)


def _yeni_sirali(kayitlar):
    ep = datetime.min.replace(tzinfo=timezone.utc)
    return sorted(kayitlar, key=lambda k: k.get("zaman") or ep, reverse=True)


@st.cache_data(ttl=120, show_spinner="Eşleşmeler hesaplanıyor...")
def _hesapla(adlar, bolgeler, secim):
    """Akış A ve B'nin tüm katmanları.
    adlar: tuple (kişinin ad varyantları) ya da None = TÜM EKİP (yalnız yönetici görünümü).
    bolgeler: tuple; secim: None (tümü) ya da bir ilçe adı."""
    ekip = adlar is None
    veri, hatalar, bilgi = _veri()
    hatalar = dict(hatalar)
    simdi = datetime.now(timezone.utc)
    talepler = es.omre_gore(veri["talep"], simdi)
    paylasim = [k for k in es.omre_gore(veri["paylasim"], simdi) if _paylasim_mi(k)]
    yat, yat_hata = _yatirim(adlar, hepsi=ekip) if ekip else _yatirim(adlar)
    if yat_hata:
        hatalar["Yatırım talebi"] = yat_hata

    if ekip:
        benim_p = veri["zeta"] + paylasim
        benim_t = [k for k in talepler if _zeta_talebi_mi(k)] + yat   # Taleplerim: yalnız Zeta danışmanlarının talepleri
        gorulen = {}
        for k in benim_p + benim_t:
            for i in k.get("ilceler") or []:
                if i:
                    gorulen.setdefault(bh.tr_kucuk(i), i)
        ilceler, ek = sorted(gorulen.values(), key=bg.tr_sira), []
        kume = None if secim is None else bg._ilce_kumesi([secim])
    else:
        benim_p = [k for k in veri["zeta"] + paylasim if bg.benim_mi(k, adlar)]
        benim_t = [k for k in talepler if _benim_talebim(k, adlar)] + yat   # yatırım talepleri sunucuda kişiye süzülü
        ilceler, ek = bg.kapsam_ilceleri(list(bolgeler), {"p": benim_p, "t": benim_t})
        kume = bg._ilce_kumesi(ilceler if secim is None else [secim])

    # Kendi kayıtların "tümü"nde ilçe süzgeciyle ELENMEZ (ilçesi belirsiz olsa da listede görünür — mevcut
    # Radarım'daki "kendi kayıtlar tüm ilçelerden" kuralı). Bir ilçe seçilince o ilçeye göre süzülür.
    kendi_hepsi = (not ekip) and secim is None
    port = list(benim_p) if kendi_hepsi else [k for k in benim_p if _bolgede(k, kume)]
    bolge = [k for k in talepler if _bolgede(k, kume)] + (list(yat) if kendi_hepsi else [k for k in yat if _bolgede(k, kume)])   # ağın tüm talepleri
    tal = list(benim_t) if kendi_hepsi else [k for k in benim_t if _bolgede(k, kume)]
    # İlçesi belirsiz kendi kayıtları listelerde görünür ama eşleşmeye girmez (ilçesiz kayıt her ilçeyle
    # "uyuşur", sahte eşleşme üretir).
    belirsiz = 0 if ekip else sum(1 for k in port if not (k["motor"].get("ilce") or "").strip()) \
        + sum(1 for k in tal if not k.get("ilceler"))

    # Ekip görünümünde hesap yükünü sınırla: en yeni kayıtlar
    kirpildi = {}
    port_h = [k for k in port if (k["motor"].get("ilce") or "").strip()]
    tal_h = [k for k in tal if k.get("ilceler")]
    bolge_h = [k for k in bolge if k.get("ilceler") or not k.get("yatirim")]
    if ekip:
        if len(port_h) > EKIP_PORT_SINIR:
            port_h, kirpildi["port"] = _yeni_sirali(port_h)[:EKIP_PORT_SINIR], len(port)
        if len(tal_h) > EKIP_TALEP_SINIR:
            tal_h, kirpildi["tal"] = _yeni_sirali(tal_h)[:EKIP_TALEP_SINIR], len(tal)

    # Benzer Startkey / FSBO: portföyleri ilçeye göre indeksle
    idx = {}
    for p in port:
        idx.setdefault((p["motor"].get("ilce") or "").strip().lower(), []).append(p)

    def _benzerleri(pazar):
        sonuc = []
        for s in pazar:
            adaylar = idx.get((s["motor"].get("ilce") or "").strip().lower())
            if not adaylar:
                continue
            esl = [p for p in adaylar if _benzer_mi(p["motor"], s["motor"])]
            if esl:
                sonuc.append((s, esl))
        return sonuc

    benz = _benzerleri(veri["startkey"])
    benz_f = _benzerleri(veri.get("fsbo", []))

    ciftler = es.portfoyden_taleplere(port_h, bolge_h, ust_n=5, min_skor=es.OLASI_ESIK)

    # ── Akış B: Taleplerim → uygun portföyler (kaynağa göre) → Eşleşmeler ──
    kaynaklar_b = {"zp": veri["zeta"], "sk": veri["startkey"], "fs": veri.get("fsbo", []), "pay": paylasim}
    uygun, ciftler_b = {}, []
    for ad, adaylar in kaynaklar_b.items():
        cf = es.talepten_portfoylere(tal_h, adaylar, ust_n=5, min_skor=es.OLASI_ESIK) if tal_h else []
        ciftler_b.extend(cf)
        gor, sira = {}, []
        for c in cf:
            kid = c["aday"]["kid"]
            if kid not in gor:
                gor[kid] = (c["aday"], [])
                sira.append(kid)
            gor[kid][1].append((c["kaynak"], c["skor"]))
        uygun[ad] = [(gor[k][0], sorted(gor[k][1], key=lambda x: -x[1])) for k in sira]
    return {
        "ekip": ekip, "ilceler": ilceler, "ek": ek, "port": port, "bolge": bolge, "benz": benz, "benz_f": benz_f,
        "ciftler": ciftler, "tal": tal, "uygun": uygun, "ciftler_b": ciftler_b,
        "kirpildi": kirpildi, "belirsiz": belirsiz,
        "hatalar": hatalar, "bayat": (bilgi or {}).get("bayat_zeta", 0), "sinir": (bilgi or {}).get("sinir", []),
    }


def _ekip_tablosu(h, kisiler):
    """Danışman bazında özet satırları (yalnız yönetici 'Tüm ekip' görünümü)."""
    ad_of = {bh.tr_kucuk(k): k for k in kisiler}
    sat = {}

    def _s(sahip):
        ad = _zeta_kisisi(sahip) or (_sahip_ayir(sahip)[0] or str(sahip).strip() or "(sahipsiz)")
        a = bh.tr_kucuk(ad) or "—"
        ad = ad_of.get(a) or ad
        return sat.setdefault(a, {"Danışman": ad, "Portföy": 0, "Talep": 0, "Güçlü (portföy→talep)": 0, "Güçlü (talep→portföy)": 0})

    for k in h["port"]:
        _s(k.get("sahip"))["Portföy"] += 1
    for k in h["tal"]:
        _s(k.get("sahip"))["Talep"] += 1
    for c in h["ciftler"]:
        if c["guclu"]:
            _s(c["kaynak"].get("sahip"))["Güçlü (portföy→talep)"] += 1
    for c in h["ciftler_b"]:
        if c["guclu"]:
            _s(c["kaynak"].get("sahip"))["Güçlü (talep→portföy)"] += 1
    return sorted(sat.values(), key=lambda r: (-(r["Portföy"] + r["Talep"]), bg.tr_sira(r["Danışman"])))


# ── GÖSTERİM YARDIMCILARI ───────────────────────────────────────────────
def _dagilim(kayitlar, kume_ilce=None):
    """{ilçe_anahtarı: [ilçe_adı, adet]} — bir kayıt birden çok ilçedeyse her birinde sayılır."""
    d = {}
    for k in kayitlar:
        for a, i in {bh.tr_kucuk(i): i for i in k.get("ilceler") or [] if i}.items():
            if kume_ilce is None or a in kume_ilce:
                e = d.setdefault(a, [i, 0])
                e[1] += 1
    return d


def _ortak_sira(dags, ref="port"):
    """Tüm kutularda AYNI ilçe sırası: önce Portföylerim'in yoğunluk sırası, sonra (portföyü olmayan ilçeler)
    tüm kutulardaki toplam yoğunluğa göre. Dönüş: (sıralı ilçe anahtarları, {anahtar: ad})."""
    toplam, ad_of = {}, {}
    for d in dags.values():
        for a, (ad, n) in d.items():
            toplam[a] = toplam.get(a, 0) + n
            ad_of.setdefault(a, ad)
    p = dags.get(ref, {})
    sira = sorted(toplam, key=lambda a: (-(p[a][1] if a in p else 0), -toplam[a], bg.tr_sira(ad_of[a])))
    return sira, ad_of


def _kutu_html(simge, baslik, sayi, dagilim, ilk=False, sira=None, ad_of=None, satir=4):
    sira = sira if sira is not None else sorted(dagilim, key=lambda a: (-dagilim[a][1], bg.tr_sira(dagilim[a][0])))
    ad_of = ad_of or {a: v[0] for a, v in dagilim.items()}
    gor = sira[:satir]
    satirlar = ""
    for a in gor:
        n = dagilim[a][1] if a in dagilim else 0
        satirlar += (f"<div class='yr-sifir'><b>0</b> {_html.escape(ad_of[a])}</div>" if n == 0
                     else f"<div><b>{n}</b> {_html.escape(ad_of[a])}</div>")
    kalan = sum(1 for a, v in dagilim.items() if a not in gor and v[1] > 0)
    if kalan > 0:
        satirlar += f"<div class='yr-ek'>+{kalan} ilçe</div>"
    if not satirlar:
        satirlar = "<div class='yr-bos'>—</div>"
    return (
        f"<div class='yr-kutu{' yr-ilk' if ilk else ''}'>"
        f"<div class='yr-sol'><div class='yr-bas'>{simge}{_html.escape(baslik)}</div><div class='yr-sayi'>{sayi}</div></div>"
        f"<div class='yr-sag'>{satirlar}</div></div>"
    )


_CSS = """
<style>
div[class*="st-key-yr_akis_"] > div[data-testid="stVerticalBlockBorderWrapper"] { border-radius: 14px; }
div[class*="st-key-yr_akis_a"] { border-top: 5px solid #1b2742; }
div[class*="st-key-yr_akis_b"] { border-top: 5px solid #2e7d6f; }
.yr-kutu { display:flex; justify-content:space-between; gap:10px; align-items:flex-start; }
.yr-sol { min-width:0; }
.yr-bas { font-size:12.5px; font-weight:700; color:#1b2742; line-height:1.25; }
.yr-sayi { font-size:24px; font-weight:800; color:#1b2742; line-height:1.15; margin-top:2px; }
.yr-ilk .yr-bas { font-size:15px; } .yr-ilk .yr-sayi { font-size:34px; }
.yr-sag { font-size:11.5px; color:#5b6478; text-align:left; line-height:1.45; border-left:1px solid #e6dfd0; padding-left:10px; min-width:84px; }
.yr-sag b { color:#1b2742; } .yr-ek, .yr-bos { color:#9a9488; } .yr-sifir, .yr-sifir b { color:#b9b4a8 !important; }
div[class*="st-key-yr_box_"] div[data-testid="stVerticalBlockBorderWrapper"] { border:1px solid #e6dfd0 !important; border-radius:12px; padding:10px 12px !important; background:#fff; }
div[class*="st-key-yr_box_es_"] div[data-testid="stVerticalBlockBorderWrapper"] { background:#fbf4e1; }
div[class*="st-key-yr_box_"][class*="_secili"] div[data-testid="stVerticalBlockBorderWrapper"] { border:2px solid #b08a2e !important; }
div[class*="st-key-yr_box_"] [data-testid="stHorizontalBlock"] { flex-wrap:nowrap !important; align-items:flex-start; }
div[class*="st-key-yr_box_"] [data-testid="stHorizontalBlock"] > div[data-testid="stColumn"] { min-width:0 !important; }
div[class*="st-key-yr_ok_"] button {
    width:32px !important; height:32px !important; min-height:32px !important; border-radius:50% !important;
    background:#f2ede0 !important; border:none !important; color:#b08a2e !important; padding:0 !important; font-weight:700 !important;
}
div[class*="st-key-yr_ok_"] button:hover { background:#e7dfc8 !important; }
.yr-ok-asagi { text-align:center; color:#b08a2e; font-weight:700; line-height:1; margin:2px 0; }
div[class*="st-key-yr_detay"] { border-top:4px solid #b08a2e; }
</style>
"""


def _acik():
    return list(st.session_state.get("yr_acik") or [])


def _sec(k):
    """Kutuya basınca listesi kutunun altında açılır; aynı kutuya tekrar basınca kapanır.
    Kutular birbirinden bağımsızdır: birini açmak diğerlerini kapatmaz."""
    acik = _acik()
    if k in acik:
        acik.remove(k)
    else:
        acik.append(k)
    st.session_state["yr_acik"] = acik


def _kapat(k):
    st.session_state["yr_acik"] = [x for x in _acik() if x != k]


def _kutu(anahtar, simge, baslik, sayi, dagilim, ilk=False, sira=None, ad_of=None):
    secili = anahtar in _acik()
    with st.container(border=True, key=f"yr_box_{anahtar}_{'secili' if secili else 'degil'}"):
        c1, c2 = st.columns([7, 1])
        with c1:
            st.markdown(_kutu_html(simge, baslik, sayi, dagilim, ilk, sira, ad_of), unsafe_allow_html=True)
        with c2:
            st.button(
                "▴" if secili else "▾", key=f"yr_ok_{anahtar}", on_click=_sec, args=(anahtar,),
                help=f"{baslik} listesini {'kapat' if secili else 'aç'}",
            )


def _tablo(kayitlar, anahtar):
    if not kayitlar:
        st.info("Bu kapsamda kayıt yok.")
        return
    simdi = datetime.now(timezone.utc)
    kolonlar = [
        {"k": "Kayıt", "label": "Kayıt", "w": 230},
        {"k": "İlan", "label": "İlan", "w": 70, "link": True},
        {"k": "İşlem", "label": "İşlem", "w": 90},
        {"k": "Fiyat", "label": "Fiyat / Bütçe", "w": 120},
        {"k": "Ayrıntı", "label": "Ayrıntı", "w": 210},
        {"k": "Kaynak", "label": "Kaynak · tarih", "w": 190},
        {"k": "Sahibi", "label": "Sahibi", "w": 120},
        {"k": "İlçe", "label": "İlçe", "w": 100},
        {"k": "Zaman", "label": "Zaman", "w": 100, "sk": True},
    ]
    satirlar = [{
        "id": k["kid"], "yeni": False, "u": bh.guvenli_url(k.get("link")),
        "t": k["zaman"].timestamp() if k.get("zaman") else 0,
        "h": {
            "Kayıt": k["baslik"], "İşlem": k.get("islem") or "—", "Fiyat": k.get("fiyat") or "—",
            "Ayrıntı": k.get("alt") or "",
            "Kaynak": f"{k.get('kaynak') or '—'} · {bh.tarih_metni(k.get('zaman'))}" if k.get("zaman") else (k.get("kaynak") or "—"),
            "Sahibi": k.get("sahip") or "—",
            "İlçe": ", ".join(i for i in k.get("ilceler") or [] if i) or "—",
            "Zaman": bh.zaman_etiketi(k.get("zaman"), simdi),
        },
    } for k in kayitlar]
    havuz_tablosu(
        kolonlar, satirlar, key=f"yr_tbl_{anahtar}", secimsiz=True,
        kart={"baslik": "Kayıt", "meta": ["Kaynak"], "alt": ["İlçe", "İşlem", "Fiyat", "Sahibi"]},
    )


def _eslesme_detay(ciftler, anahtar, sec):
    duzey = st.radio("Göster", ["Sadece güçlü (%70+)", "Güçlü + olası (%40+)"], key=f"yr_duzey_{sec}", horizontal=True)
    gorunen = [c for c in ciftler if c["guclu"]] if duzey.startswith("Sadece") else ciftler
    if not gorunen:
        st.info("Bu filtrelerde eşleşme yok.")
        return
    gruplar = es.gruplari_kur(gorunen)
    n = st.session_state.get(f"yr_n_{anahtar}", GRUP_SAYFA)
    secilen = gruplar[:n]
    imza = "|".join([str(len(gruplar)), str(n), duzey, anahtar] + [g["kaynak"]["baslik"][:20] + str(g["en_iyi"]) for g in secilen[:5]])
    eslesme_kartlari(
        [{"kaynak": g["kaynak"], "adaylar": g["adaylar"]} for g in secilen],
        es.kart_css(), imza, key=f"yr_esk_{anahtar}",
    )
    st.caption(f"{len(secilen)} / {len(gruplar)} portföy gösteriliyor")
    if len(gruplar) > n and st.button(f"Daha fazla göster (+{GRUP_SAYFA})", key=f"yr_daha_{sec}"):
        st.session_state[f"yr_n_{anahtar}"] = n + GRUP_SAYFA
        st.rerun()


def _benzer_tablo(benz):
    """Benzer Startkey ilanları, hangi portföyüme benzedikleri ayrıntıda."""
    kayitlar = []
    for s, esl in benz:
        k = dict(s)
        k["alt"] = ("Benzediği portföyün: " + " · ".join(p["baslik"][:40] for p in esl[:2])
                    + (f" (+{len(esl) - 2})" if len(esl) > 2 else ""))
        kayitlar.append(k)
    return kayitlar


def _uygun_tablo(uygun):
    """Taleplerime uyan portföyler; hangi talebime kaç puanla uyduğu ayrıntıda."""
    kayitlar = []
    for p, eslesenler in uygun:
        k = dict(p)
        k["alt"] = ("Uyduğu talebin: " + " · ".join(f"{t['baslik'][:36]} (%{sk})" for t, sk in eslesenler[:2])
                    + (f" (+{len(eslesenler) - 2})" if len(eslesenler) > 2 else ""))
        kayitlar.append(k)
    return kayitlar


def _detay(sec, baslik, h, secim_ad, kullanici):
    with st.container(border=True, key=f"yr_detay_{sec}"):
        t1, t2 = st.columns([6, 2])
        with t1:
            st.markdown(f"**{_html.escape(baslik)} · {_html.escape(secim_ad)}**")
        with t2:
            st.button("✕ Kapat", key=f"yr_kapat_{sec}", on_click=_kapat, args=(sec,))
        anahtar = f"{kullanici}_{sec}_{secim_ad}"
        if sec == "port":
            _tablo(h["port"], anahtar)
        elif sec == "bolge":
            _tablo(h["bolge"], anahtar)
        elif sec == "benz":
            _tablo(_benzer_tablo(h["benz"]), anahtar)
        elif sec == "benzf":
            _tablo(_benzer_tablo(h["benz_f"]), anahtar)
        elif sec == "es":
            _eslesme_detay(h["ciftler"], anahtar, sec)
        elif sec == "tal":
            _tablo(h["tal"], anahtar)
        elif sec in ("zp", "sk", "fs", "pay"):
            _tablo(_uygun_tablo(h["uygun"][sec]), anahtar)
        else:
            _eslesme_detay(h["ciftler_b"], anahtar, sec)


# (anahtar, üst başlık, alt başlık) — tek başlık altında iki sütun olanlar alt başlıklıdır
PIYASA_KOLONLAR = (
    ("fsbo", "FSBO ilanı", ""), ("sk", "Startkey ilanı", ""), ("zp", "Zeta portföyleri", ""),
    ("kp_sk", "Kapalı portföyler", "Startkey"), ("kp_z", "Kapalı portföyler", "Zeta"),
    ("at_sk", "Alıcı talepleri", "Startkey"), ("at_z", "Alıcı talepleri", "Zeta"),
)
PIYASA_DONEMLER = (("toplam", "Toplam (aktif)", None), ("ay", "Son 30 gün (yeni)", 30), ("hafta", "Son 7 gün (yeni)", 7))
BELIRSIZ_ILCE = "(ilçe belirsiz)"


@st.cache_data(ttl=300, show_spinner="Startkey / FSBO ilan sayıları okunuyor...")
def _pazar_hafif():
    """Startkey ve FSBO aktif ilanları için yalnız (ilçeler, ilk görülme) — SINIRSIZ ve hafif sorgu.
    Dönüş: {'startkey': [(ilçeler, zaman)], 'fsbo': [...]}"""
    simdi = datetime.now(timezone.utc)
    kesim_iso = (simdi - timedelta(days=3650)).isoformat()
    sonuc = {"startkey": [], "mulk_sahibi": []}
    bas, boy = 0, 1000
    while bas < 300000:
        satirlar = (
            bh.supabase.table("izmir_pazar_ilanlar").select("marka,ilce,ilk_gorulme_tarihi")
            .eq("aktif", True).in_("marka", list(sonuc)).gte("ilk_gorulme_tarihi", kesim_iso)
            .order("id", desc=True).range(bas, bas + boy - 1).execute().data or []
        )
        for v in satirlar:
            if v.get("marka") in sonuc:
                sonuc[v["marka"]].append((bh.ilceleri_ayikla(v.get("ilce")), bh._zaman(v.get("ilk_gorulme_tarihi"))))
        if len(satirlar) < boy:
            break
        bas += boy
    return {"startkey": sonuc["startkey"], "fsbo": sonuc["mulk_sahibi"]}


def _ogeler(kayitlar):
    return [(k.get("ilceler") or [], k.get("zaman")) for k in kayitlar]


def _donemde(ogeler, gun, simdi):
    if gun is None:
        return ogeler
    sinir = simdi - timedelta(days=gun)
    return [o for o in ogeler if o[1] and o[1] >= sinir]


def _ilce_say(ogeler):
    """{anahtar: [ad, adet]} — çok ilçeli kayıt her ilçede sayılır; ilçesiz olanlar BELIRSIZ_ILCE altında."""
    d = {}
    for ilceler, _ in ogeler:
        gecerli = {bh.tr_kucuk(i): i for i in ilceler if i}
        if not gecerli:
            gecerli = {"": BELIRSIZ_ILCE}
        for a, i in gecerli.items():
            e = d.setdefault(a, [i, 0])
            e[1] += 1
    return d


@st.cache_data(ttl=300, show_spinner="Startkey piyasa özeti hazırlanıyor...")
def _piyasa():
    """Yönetici 'Tüm Startkey' özeti: ilçe bazında arz (ilan/portföy) ve talep sayıları; Toplam / Son 30 gün / Son 7 gün.
    Eşleşme hesabı YOK. Kapalı portföy = ofis içi paylaşım: 'Zeta' = Radar'dan elle girilen (kaynak zeta/ofis),
    'Startkey' = Startkey mail sistemindeki diğer paylaşımlar. Alıcı talebi: 'Zeta' = Zeta danışmanlarının talepleri
    + yatırım talepleri, 'Startkey' = ağın diğer talepleri.
    Toplam: aktif stok (talep/paylaşım kiralık 14, satılık 45 gün); Son 30/7 gün: o sürede gelen yeni kayıtlar.
    İlçe sırası: Toplam dönemindeki FSBO ilanı adedi (pazar büyüklüğü), çoktan aza; tüm dönemlerde aynı sıra.
    Toplam dönemde Zeta portföyü (Zeta portföyleri ya da Zeta kapalı portföyü) olan ilçeler ayrı satırda; Zeta portföyü olmayanlar
    tek 'Diğer ilçeler' satırında toplanır (ayrıntısı açılır); ilçesi belirsiz kayıtlar en altta.
    Satır: {'ilce', 'grup': 'ana'|'diger'|'belirsiz', kolonlar...}. Dönüş: ({dönem: (satırlar, toplamlar)}, hatalar, tarih_metni)."""
    veri, hatalar, _ = _veri_temel()
    hatalar = dict(hatalar)
    simdi = datetime.now(timezone.utc)
    yat, yat_hata = _yatirim(None, hepsi=True)
    if yat_hata:
        hatalar["Yatırım talebi"] = yat_hata
    try:
        pazar = _pazar_hafif()
    except Exception as e:      # noqa: BLE001
        pazar = {"startkey": [], "fsbo": []}
        hatalar["Startkey/FSBO ilanı"] = str(e)

    sonuc = {}
    for anahtar, _ad, gun in PIYASA_DONEMLER:
        if gun is None:      # aktif stok: talep/paylaşım türüne göre ömür süzgeci
            talepler = es.omre_gore(veri["talep"], simdi)
            paylasim = es.omre_gore(veri["paylasim"], simdi)
        else:                # o sürede gelen yeni kayıtlar (ömür süzgeci yok)
            talepler, paylasim = veri["talep"], veri["paylasim"]
        zeta_ids = {k["kid"] for k in talepler if _zeta_talebi_mi(k)}
        kaynaklar = {
            "fsbo": pazar["fsbo"], "sk": pazar["startkey"], "zp": _ogeler(veri["zeta"]),
            "kp_sk": _ogeler([k for k in paylasim if not _paylasim_mi(k)]),
            "kp_z": _ogeler([k for k in paylasim if _paylasim_mi(k)]),
            "at_sk": _ogeler([k for k in talepler if k["kid"] not in zeta_ids]),
            "at_z": _ogeler([k for k in talepler if k["kid"] in zeta_ids] + yat),
        }
        kaynaklar = {ad: _donemde(o if isinstance(o, list) and (not o or isinstance(o[0], tuple)) else _ogeler(o), gun, simdi)
                     for ad, o in kaynaklar.items()}
        dags = {ad: _ilce_say(o) for ad, o in kaynaklar.items()}
        ad_of, anahtarlar = {}, set()
        for d in dags.values():
            for a, (ad, _) in d.items():
                ad_of.setdefault(a, ad)
                anahtarlar.add(a)
        ham = {}
        for a in anahtarlar:
            r = {"ilce": ad_of[a], "_a": a}
            for kol, *_ in PIYASA_KOLONLAR:
                r[kol] = dags[kol][a][1] if a in dags[kol] else 0
            ham[a] = r
        sonuc[anahtar] = (ham, {kol: len(o) for kol, o in kaynaklar.items()})
    # sıra ve 'diğer' kümesi Toplam dönemine göre bir kez belirlenir, her dönemde aynı uygulanır
    toplam_ham = sonuc[PIYASA_DONEMLER[0][0]][0]
    sirali = sorted((a for a in toplam_ham if a != ""), key=lambda a: (-toplam_ham[a]["fsbo"], -toplam_ham[a]["sk"], bg.tr_sira(toplam_ham[a]["ilce"])))
    zetali = lambda a: toplam_ham[a]["zp"] + toplam_ham[a]["kp_z"] > 0
    ana = [a for a in sirali if zetali(a)]
    diger = [a for a in sirali if not zetali(a)]
    for anahtar in list(sonuc):
        ham, toplam = sonuc[anahtar]
        bos = lambda ad, a: {"ilce": ad, "_a": a, **{kol: 0 for kol, *_ in PIYASA_KOLONLAR}}
        satirlar = [dict(ham.get(a) or bos(toplam_ham[a]["ilce"], a), grup="ana") for a in ana]
        if diger:
            alt = [dict(ham.get(a) or bos(toplam_ham[a]["ilce"], a), grup="detay") for a in diger]
            ozet = {"ilce": f"Diğer ilçeler ({len(diger)})", "grup": "diger", "n": len(diger),
                    **{kol: sum(r[kol] for r in alt) for kol, *_ in PIYASA_KOLONLAR}}
            satirlar += [ozet] + alt
        if "" in ham:
            satirlar.append(dict(ham[""], grup="belirsiz"))
        sonuc[anahtar] = (satirlar, toplam)
    return sonuc, hatalar, bh.tarih_metni(simdi)


_PIYASA_CSS = """
.yr-pz-kap { overflow:auto; max-height:640px; border:1px solid #dcdfe4; border-radius:12px; background:#fff; }
table.yr-pz { border-collapse:separate; border-spacing:0; width:100%; font-size:13.5px; }
table.yr-pz th, table.yr-pz td { padding:7px 10px; text-align:right; border-bottom:1px solid #eceef1; white-space:nowrap; }
table.yr-pz th { background:#e9ecef; color:#2b3340; font-weight:700; position:sticky; top:0; z-index:2; text-align:center; }
table.yr-pz thead tr:nth-child(2) th { top:34px; font-weight:600; font-size:12.5px; background:#f1f3f5; }
table.yr-pz th.yr-g { border-left:1px solid #dcdfe4; }
table.yr-pz td:first-child, table.yr-pz th:first-child { text-align:left; position:sticky; left:0; background:#fff; z-index:1; }
table.yr-pz th:first-child { background:#e9ecef; z-index:3; }
table.yr-pz td.yr-g { border-left:1px solid #eceef1; }
table.yr-pz td.yr-0 { color:#b9bec7; }
table.yr-pz tr.yr-top td { font-weight:800; background:#e9ecef; }
table.yr-pz tr.yr-diger td { background:#f6f7f9; font-weight:700; }
table.yr-pz tr.yr-diger td:first-child { background:#f6f7f9; }
table.yr-pz tr.yr-diger.yr-tikla td { cursor:pointer; }
table.yr-pz tr.yr-det { display:none; }
table.yr-pz.ac tr.yr-det { display:table-row; }
table.yr-pz tr.yr-det td:first-child { padding-left:26px; color:#4b5563; }
table.yr-pz tr.yr-bel td { color:#6b7280; font-style:italic; }
"""


def _baslik_yapisi():
    """[(metin, kolon_sayisi, satir_sayisi)] üst başlıklar, [alt başlıklar]."""
    ust, alt, kolonlar, i = [], [], list(PIYASA_KOLONLAR), 0
    while i < len(kolonlar):
        _, u, a = kolonlar[i]
        if a:
            n = 1
            while i + n < len(kolonlar) and kolonlar[i + n][1] == u:
                n += 1
            ust.append((u, n, 1))
            alt.extend(x[2] for x in kolonlar[i:i + n])
            i += n
        else:
            ust.append((u, 1, 2))
            i += 1
    return ust, alt


def _piyasa_html(satirlar, toplam, acik=False, tikla=False):
    """acik: 'Diğer ilçeler' ayrıntı satırları görünür; tikla: satıra tıklayınca açılır (yalnız indirilen HTML'de, JS gerekir)."""
    ust, alt = _baslik_yapisi()
    h1 = "<th rowspan='2'>İlçe</th>" + "".join(
        f"<th colspan='{n}' class='yr-g'>{_html.escape(u)}</th>" if r == 1 else f"<th rowspan='2' class='yr-g'>{_html.escape(u)}</th>"
        for u, n, r in ust)
    h2 = "".join(f"<th class='yr-g'>{_html.escape(a)}</th>" for a in alt)
    govde = ""
    for r in satirlar:
        grup = r.get("grup", "ana")
        sinif, ad = {"diger": ("yr-diger", ""), "detay": ("yr-det", ""), "belirsiz": ("yr-bel", "")}.get(grup, ("", ""))
        etiket = _html.escape(r["ilce"])
        ek = ""
        if grup == "diger":
            etiket = "▸ " + etiket
            if tikla:
                sinif += " yr-tikla"
                ek = " onclick=\"this.closest('table').classList.toggle('ac')\" title='Ayrıntı için tıkla'"
        govde += f"<tr class='{sinif}'{ek}><td>{etiket}</td>" + "".join(
            f"<td class='yr-g{' yr-0' if not r[k] else ''}'>{r[k]}</td>" for k, *_ in PIYASA_KOLONLAR) + "</tr>"
    govde += "<tr class='yr-top'><td>TOPLAM (kayıt)</td>" + "".join(
        f"<td class='yr-g'>{toplam[k]}</td>" for k, *_ in PIYASA_KOLONLAR) + "</tr>"
    return (f"<div class='yr-pz-kap'><table class='yr-pz{' ac' if acik else ''}'><thead><tr>{h1}</tr><tr>{h2}</tr></thead>"
            f"<tbody>{govde}</tbody></table></div>")


_PIYASA_NOT_METNI = (
    "FSBO / Startkey ilanı: aktif ilanlar · Zeta portföyleri: Revy'de görünen ofis stoku · Kapalı portföyler: ofis içi paylaşımlar "
    "(Zeta: Radar'dan elle girilen; Startkey: mail sistemindeki diğerleri) · Alıcı talepleri (Zeta: Zeta danışmanlarının talepleri ve "
    "formdan gelen yatırım talepleri; Startkey: ağın diğer talepleri). Toplam: aktif stok (talep ve paylaşımlar kiralıkta 14, satılıkta 45 gün); "
    "Son 30 / 7 gün: o sürede gelen yeni kayıtlar. Zeta portföyü olan ilçeler Toplam dönemindeki FSBO ilanı adedine göre (pazar büyüklüğü) sıralıdır; Zeta portföyü olmayan ilçeler \"Diğer ilçeler\" satırında toplanır. Birden çok ilçe isteyen kayıt her ilçede sayılır; son satır kayıt sayısıdır."
)


def _piyasa_not():
    return _PIYASA_NOT_METNI


def _piyasa_excel(sonuc, tarih):
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    wb, ilk = Workbook(), True
    ust, alt = _baslik_yapisi()
    krem, koyu = PatternFill("solid", fgColor="E9ECEF"), Font(bold=True, color="2B3340")
    acik_gri = PatternFill("solid", fgColor="F6F7F9")
    ince = Side(style="thin", color="DCDFE4")
    for anahtar, ad, _ in PIYASA_DONEMLER:
        satirlar, toplam = sonuc[anahtar]
        ws = wb.active if ilk else wb.create_sheet()
        ilk = False
        ws.title = ad.split(" (")[0]
        ws["A1"] = f"Tüm Startkey — ilçe bazında piyasa özeti · {ad} · {tarih}"
        ws["A1"].font = Font(bold=True, size=13, color="2B3340")
        ws.sheet_properties.outlinePr.summaryBelow = False
        ws.cell(row=3, column=1, value="İlçe")
        ws.merge_cells(start_row=3, start_column=1, end_row=4, end_column=1)
        c, alt_i = 2, 0
        for u, n, r in ust:
            ws.cell(row=3, column=c, value=u)
            if r == 2:
                ws.merge_cells(start_row=3, start_column=c, end_row=4, end_column=c)
            else:
                if n > 1:
                    ws.merge_cells(start_row=3, start_column=c, end_row=3, end_column=c + n - 1)
                for j in range(n):
                    ws.cell(row=4, column=c + j, value=alt[alt_i])
                    alt_i += 1
            c += n
        for rr in (3, 4):
            for cc in range(1, 2 + len(PIYASA_KOLONLAR)):
                h = ws.cell(row=rr, column=cc)
                h.fill, h.font = krem, koyu
                h.alignment = Alignment(horizontal="center", vertical="center")
                h.border = Border(bottom=ince)
        for i, r in enumerate(satirlar):
            sr = 5 + i
            grup = r.get("grup", "ana")
            ws.cell(row=sr, column=1, value=("    " if grup == "detay" else "") + r["ilce"])
            for j, (k, *_) in enumerate(PIYASA_KOLONLAR):
                ws.cell(row=sr, column=2 + j, value=r[k])
            if grup == "diger":
                for cc in range(1, 2 + len(PIYASA_KOLONLAR)):
                    ws.cell(row=sr, column=cc).font = Font(bold=True)
                    ws.cell(row=sr, column=cc).fill = acik_gri
            elif grup == "detay":      # Excel'de + işaretiyle açılıp kapanan gruplama (varsayılan kapalı)
                ws.row_dimensions[sr].outline_level = 1
                ws.row_dimensions[sr].hidden = True
                ws.cell(row=sr, column=1).font = Font(color="4B5563")
            elif grup == "belirsiz":
                for cc in range(1, 2 + len(PIYASA_KOLONLAR)):
                    ws.cell(row=sr, column=cc).font = Font(italic=True, color="6B7280")
        son = 5 + len(satirlar)
        ws.cell(row=son, column=1, value="TOPLAM (kayıt)")
        for j, (k, *_) in enumerate(PIYASA_KOLONLAR):
            ws.cell(row=son, column=2 + j, value=toplam[k])
        for cc in range(1, 2 + len(PIYASA_KOLONLAR)):
            ws.cell(row=son, column=cc).font = Font(bold=True, color="2B3340")
            ws.cell(row=son, column=cc).fill = krem
        ws.column_dimensions["A"].width = 22
        for cc in range(2, 2 + len(PIYASA_KOLONLAR)):
            ws.column_dimensions[get_column_letter(cc)].width = 16
        ws.freeze_panes = "B5"
        ws.cell(row=son + 2, column=1, value=_piyasa_not())
        ws.cell(row=son + 2, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=son + 2, start_column=1, end_row=son + 2, end_column=1 + len(PIYASA_KOLONLAR))
        ws.row_dimensions[son + 2].height = 90
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _piyasa_dosya_html(sonuc, tarih):
    """Tek dosyalık, telefonda/WhatsApp'ta paylaşılabilir sayfa: üç dönem sekme olarak."""
    sekmeler = "".join(
        f"<button class='sk{' on' if i == 0 else ''}' onclick=\"goster('{a}',this)\">{_html.escape(ad)}</button>"
        for i, (a, ad, _) in enumerate(PIYASA_DONEMLER))
    bolumler = "".join(
        f"<section id='{a}' style='display:{'block' if i == 0 else 'none'}'>{_piyasa_html(*sonuc[a], tikla=True)}</section>"
        for i, (a, ad, _) in enumerate(PIYASA_DONEMLER))
    return (
        "<!doctype html><html lang='tr'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
        "<title>Tüm Startkey — piyasa özeti</title><style>"
        "body{font-family:Segoe UI,Arial,sans-serif;background:#f4f5f7;color:#2b3340;margin:0;padding:16px}"
        "h1{font-size:18px;margin:0 0 4px}.alt{color:#6b7385;font-size:13px;margin-bottom:12px}"
        ".sk{border:1px solid #dcdfe4;background:#fff;border-radius:999px;padding:7px 14px;margin:0 6px 10px 0;font-size:13px;cursor:pointer}"
        ".sk.on{background:#2b3340;color:#fff;border-color:#2b3340}.not{color:#6b7385;font-size:12px;line-height:1.5;margin-top:12px}"
        + _PIYASA_CSS +
        "</style></head><body><h1>Tüm Startkey — ilçe bazında piyasa özeti</h1>"
        f"<div class='alt'>Startkey Zeta · {tarih}</div><div>{sekmeler}</div>{bolumler}"
        f"<div class='not'>{_html.escape(_piyasa_not())}</div>"
        "<script>function goster(id,b){document.querySelectorAll('section').forEach(function(s){s.style.display='none'});"
        "document.getElementById(id).style.display='block';document.querySelectorAll('.sk').forEach(function(x){x.classList.remove('on')});"
        "b.classList.add('on')}</script></body></html>"
    ).encode("utf-8")


def _piyasa_ozeti():
    c1, c2 = st.columns([5, 1])
    with c1:
        st.markdown("**🌐 Tüm Startkey — ilçe bazında piyasa özeti**")
    with c2:
        if st.button("↻ Yenile", key="yr_piyasa_yenile"):
            _veri_temizle()
            st.rerun()
    acik = st.toggle("Diğer ilçeleri aç (Zeta portföyü olmayanlar)", key="yr_piyasa_diger")
    sonuc, hatalar, tarih = _piyasa()
    for kaynak, msj in hatalar.items():
        st.warning(f"{kaynak} okunamadı, dahil edilmedi: {msj}")
    donem = st.radio("Dönem", [ad for _, ad, _ in PIYASA_DONEMLER], horizontal=True, key="yr_piyasa_donem")
    anahtar = next(a for a, ad, _ in PIYASA_DONEMLER if ad == donem)
    satirlar, toplam = sonuc[anahtar]
    if not satirlar:
        st.info("Bu dönemde kayıt yok.")
    else:
        st.markdown(f"<style>{_PIYASA_CSS}</style>" + _piyasa_html(satirlar, toplam, acik=acik), unsafe_allow_html=True)
    st.caption(_piyasa_not())
    d1, d2, _ = st.columns([2, 2, 3])
    gun = tarih.replace(".", "")
    with d1:
        st.download_button("⬇ Excel (3 dönem)", data=_piyasa_excel(sonuc, tarih), file_name=f"startkey_piyasa_ozeti_{gun}.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="yr_piyasa_xlsx")
    with d2:
        st.download_button("⬇ Paylaşılabilir sayfa (HTML)", data=_piyasa_dosya_html(sonuc, tarih), file_name=f"startkey_piyasa_ozeti_{gun}.html",
                           mime="text/html", key="yr_piyasa_html")


def _kisi_ac(ad):
    st.session_state["yr_kisi"] = ad


def _ekip_bolumu(h, kisiler):
    import pandas as pd
    st.markdown("**👥 Danışman dağılımı**")
    satirlar = _ekip_tablosu(h, kisiler)
    if not satirlar:
        st.info("Bu kapsamda kayıt yok.")
        return
    toplam = {"Danışman": "TOPLAM"}
    for kol in ("Portföy", "Talep", "Güçlü (portföy→talep)", "Güçlü (talep→portföy)"):
        toplam[kol] = sum(r[kol] for r in satirlar)
    st.dataframe(pd.DataFrame(satirlar + [toplam]), hide_index=True, use_container_width=True)
    adlar_listesi = [r["Danışman"] for r in satirlar if r["Danışman"] in kisiler]
    if adlar_listesi:
        c1, c2 = st.columns([4, 1])
        with c1:
            sec = st.selectbox("Danışmanın Radarım'ını aç", adlar_listesi, key="yr_ekip_ac")
        with c2:
            st.write("")
            st.button("Aç →", key="yr_ekip_ac_btn", on_click=_kisi_ac, args=(sec,))


# ── GİRİŞ NOKTASI ───────────────────────────────────────────────────────
def render_yeni_radarim():
    st.markdown(_CSS, unsafe_allow_html=True)
    kullanici = su_anki_danisman()
    adlar = tuple(su_anki_danisman_varyantlari() or [kullanici])

    # Yönetici: başka bir danışmanın ya da TÜM EKİBİN Radarım'ını görüntüleyebilir.
    # Danışmanlar yalnız kendi verisini görür (yatırım talepleri de yalnız atanan danışmana + yöneticiye).
    ekip, kisiler = False, []
    if _yonetici():
        try:
            from core.personel_manager import load_personel_listesi
            _df = load_personel_listesi()
            kisiler = sorted({str(a).strip() for a in _df["ad_soyad"].dropna() if str(a).strip()}, key=bg.tr_sira)
        except Exception:      # noqa: BLE001
            kisiler = []
        secilen_kisi = st.selectbox("Görünüm (yönetici)", ["Kendim", EKIP, PIYASA] + kisiler, key="yr_kisi")
        if secilen_kisi == PIYASA:
            _piyasa_ozeti()
            return
        if secilen_kisi == EKIP:
            ekip, kullanici, adlar = True, "__ekip__", None
            st.info("👥 Tüm ekibin verisini görüntülüyorsun (yönetici görünümü).")
        elif secilen_kisi != "Kendim":
            kullanici, adlar = secilen_kisi, (secilen_kisi,)
            st.info(f"👁 {secilen_kisi} adına görüntülüyorsun.")

    bolgeler = () if ekip else tuple(_bolgeler(kullanici))
    ilk = _hesapla(adlar, bolgeler, None)
    ilceler = ilk["ilceler"]
    tum_ad = "Tüm ilçeler" if ekip else TUM

    # İlçe seçicide her ilçenin portföy ve talep sayısı görünür; ilçeler Portföylerim yoğunluğuna göre sıralı.
    d_port, d_tal = _dagilim(ilk["port"], None), _dagilim(ilk["tal"], None)

    def _say(d, i):
        return d[bh.tr_kucuk(i)][1] if bh.tr_kucuk(i) in d else 0

    ilceler_sirali = sorted(ilceler, key=lambda i: (-_say(d_port, i), -_say(d_tal, i), bg.tr_sira(i)))

    def _etiket(i):
        if i == tum_ad:
            return f"{tum_ad} · {len(ilk['port'])} portföy · {len(ilk['tal'])} talep"
        return f"{i}{' ◦' if i in ilk['ek'] else ''} · {_say(d_port, i)} portföy · {_say(d_tal, i)} talep"

    c1, c2 = st.columns([5, 1])
    with c1:
        secenek = [tum_ad] + ilceler_sirali
        secim_ad = st.selectbox(
            "📍 İlçe seç" if ekip else "📍 Uzmanlık bölgesi seç", secenek, key=f"yr_ilce_{kullanici}",
            format_func=_etiket,
        )
    with c2:
        st.write("")
        if st.button("↻ Yenile", key="yr_yenile"):
            _veri_temizle()
            st.rerun()
    if ilk["ek"]:
        st.caption("◦ Uzmanlık bölgen değil; kendi kaydın olduğu için eklendi.")
    if not bolgeler and not ekip:
        st.info("Uzmanlık bölgen seçili değil. Uzmanlık Bölgelerim'den en fazla 5 ilçe seçersen o ilçelerdeki "
                "ekip kayıtları da burada toplanır.")
    for kaynak, msj in ilk["hatalar"].items():
        st.warning(f"{kaynak} okunamadı, dahil edilmedi: {msj}")

    secim = None if secim_ad == tum_ad else secim_ad
    h = _hesapla(adlar, bolgeler, secim)
    guclu = [c for c in h["ciftler"] if c["guclu"]]
    benz_kayit = [s for s, _ in h["benz"]]
    benz_f_kayit = [s for s, _ in h["benz_f"]]
    kume = None if secim is None else {bh.tr_kucuk(secim)}

    # ── Akış A ── (liste, basılan kutunun hemen altında açılır)
    adlar_ = {"port": "Portföylerim", "bolge": "Alıcı talepleri", "benz": "Benzer Startkey portföyleri", "benzf": "Benzer FSBO ilanları", "es": "Eşleşmeler",
              "tal": "Taleplerim", "zp": "Zeta portföyleri", "sk": "Startkey portföyleri",
              "fs": "FSBO ilanları", "pay": "Portföy paylaşımları", "es_b": "Eşleşmeler (taleplerim)"}
    acik = [k for k in _acik() if k in adlar_]

    def _altinda(anahtar):
        if anahtar in acik:
            _detay(anahtar, adlar_[anahtar], h, secim_ad, kullanici)

    guclu_b = [c for c in h["ciftler_b"] if c["guclu"]]
    uygun_k = {ad: [p for p, _ in h["uygun"][ad]] for ad in ("zp", "sk", "fs", "pay")}
    # Tüm kutularda aynı ilçe sırası (Portföylerim'in yoğunluk sırası) — ilçe bazında kuşbakışı karşılaştırma.
    dags = {
        "port": _dagilim(h["port"], kume), "bolge": _dagilim(h["bolge"], kume),
        "benz": _dagilim(benz_kayit, kume), "benzf": _dagilim(benz_f_kayit, kume),
        "es": _dagilim([c["kaynak"] for c in guclu], kume),
        "tal": _dagilim(h["tal"], kume),
        **{ad: _dagilim(v, kume) for ad, v in uygun_k.items()},
        "es_b": _dagilim([c["kaynak"] for c in guclu_b], kume),
    }
    sira, ad_of = _ortak_sira(dags)

    with st.container(border=True, key="yr_akis_a"):
        _kutu("port", "🏠 ", adlar_["port"], len(h["port"]), dags["port"], ilk=True, sira=sira, ad_of=ad_of)
        _altinda("port")
        st.markdown("<div class='yr-ok-asagi'>↓</div>", unsafe_allow_html=True)
        _kutu("bolge", "", adlar_["bolge"], len(h["bolge"]), dags["bolge"], sira=sira, ad_of=ad_of)
        _altinda("bolge")
        _kutu("benz", "", adlar_["benz"], len(benz_kayit), dags["benz"], sira=sira, ad_of=ad_of)
        _altinda("benz")
        _kutu("benzf", "", adlar_["benzf"], len(benz_f_kayit), dags["benzf"], sira=sira, ad_of=ad_of)
        _altinda("benzf")
        st.markdown("<div class='yr-ok-asagi'>↓</div>", unsafe_allow_html=True)
        _kutu("es", "", adlar_["es"], len(guclu), dags["es"], sira=sira, ad_of=ad_of)
        _altinda("es")
    st.markdown("<div style='height:8px'></div>", unsafe_allow_html=True)

    with st.container(border=True, key="yr_akis_b"):
        _kutu("tal", "📋 ", adlar_["tal"], len(h["tal"]), dags["tal"], ilk=True, sira=sira, ad_of=ad_of)
        _altinda("tal")
        st.markdown("<div class='yr-ok-asagi'>↓</div>", unsafe_allow_html=True)
        for ad in ("zp", "sk", "fs", "pay"):
            _kutu(ad, "", adlar_[ad], len(uygun_k[ad]), dags[ad], sira=sira, ad_of=ad_of)
            _altinda(ad)
        st.markdown("<div class='yr-ok-asagi'>↓</div>", unsafe_allow_html=True)
        _kutu("es_b", "", "Eşleşmeler", len(guclu_b), dags["es_b"], sira=sira, ad_of=ad_of)
        _altinda("es_b")
    st.caption("Taleplerim akışında portföy kutuları, taleplerinden en az biriyle %40+ uyuşan portföyleri sayar; "
               "Eşleşmeler yalnızca güçlü (%70+) çiftlerdir.")
    if not acik:
        st.caption("Kutulardaki ▾ düğmesine basınca ilgili liste kutunun altında açılır; istediğin kadarını aynı anda açık tutabilirsin.")
    st.caption("💼 Yatırım talepleri Alıcı talepleri / Taleplerim içinde gösterilir; her danışman yalnız kendisine "
               "atananı görür, yönetici hepsini.")
    if h["belirsiz"]:
        st.caption(f"ℹ {h['belirsiz']} kaydının ilçesi belirsiz: sayılarda ve listelerde görünür (yalnız 'tümü' seçiliyken), "
                   "ilçe bilgisi olmadığı için eşleşmeye girmez.")
    if h["kirpildi"]:
        st.caption("⏱ Ekip görünümünde eşleşmeler en yeni "
                   + (f"{EKIP_PORT_SINIR} portföy (toplam {h['kirpildi']['port']})" if "port" in h["kirpildi"] else "")
                   + (" ve " if len(h["kirpildi"]) == 2 else "")
                   + (f"{EKIP_TALEP_SINIR} talep (toplam {h['kirpildi']['tal']})" if "tal" in h["kirpildi"] else "")
                   + " için hesaplandı; bir ilçe seçerek tümünü kapsayabilirsin.")
    if ekip:
        _ekip_bolumu(h, kisiler)
    if h["sinir"]:
        st.caption("⚠ Şu ilçelerde ilan sayısı okuma sınırına (3000) takıldı, en yeni 3000'i alındı: " + ", ".join(h["sinir"]))
    if h["bayat"]:
        st.caption(f"⚠ {h['bayat']} Zeta ilanı son günlerde Revy'de görünmediği için (yayından kalkmış olabilir) "
                   "eşleştirmeye alınmadı.")
